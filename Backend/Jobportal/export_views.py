from openpyxl import Workbook
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from .models import ApplyForm
from rest_framework.permissions import AllowAny

class ExportApplicationExcelAPIView(APIView):
    permission_classes = [IsAuthenticated]
    permission_classes = [AllowAny]

    def get(self, request):

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Applications"

        headers = [
            "ID",
            "Name",
            "Email",
            "Phone",
            "Current Location",
            "LinkedIn",
            "Portfolio",
            "Experience",
            "Designation",
            "Company",
            "Notice Period",
            "Expected Salary",
            "Status",
            "Applied On",
        ]

        for col_num, header in enumerate(headers, 1):
            sheet.cell(row=1, column=col_num).value = header

        applications = ApplyForm.objects.all().order_by("-created_at")

        row = 2

        for app in applications:
            sheet.cell(row=row, column=1).value = app.id
            sheet.cell(row=row, column=2).value = app.full_name
            sheet.cell(row=row, column=3).value = app.email
            sheet.cell(row=row, column=4).value = app.phone
            sheet.cell(row=row, column=5).value = app.current_location
            sheet.cell(row=row, column=6).value = app.linkedin
            sheet.cell(row=row, column=7).value = app.portfolio
            sheet.cell(row=row, column=8).value = app.total_experience
            sheet.cell(row=row, column=9).value = app.current_role
            sheet.cell(row=row, column=10).value = app.current_company
            sheet.cell(row=row, column=11).value = app.notice_period
            sheet.cell(row=row, column=12).value = float(app.expected_salary)
            sheet.cell(row=row, column=13).value = app.status
            sheet.cell(row=row, column=14).value = app.created_at.strftime("%d-%m-%Y %H:%M")

            row += 1

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        response["Content-Disposition"] = 'attachment; filename="Applications.xlsx"'

        workbook.save(response)

        return response